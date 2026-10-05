#!/usr/bin/env python3
"""Construit les lots de 120 questions (nomenclature de l'examen AMF) et les données du site.

Entrées : donnees/questions.json (base fusionnée), manuel/*.md (intitulés des thèmes et sous-thèmes)
Sorties : docs/data/lots-data.js            (données du site)
          donnees/lots/Lots_120_questions.xlsx (un onglet par lot + sommaire)
          donnees/lots/lots_120_questions.csv

Principe : pour chaque sous-thème, les questions sont mélangées une fois (graine fixe), puis
distribuées en tourniquet, n questions par lot (n = nombre de questions du sous-thème à l'examen).
Les 13 premiers lots n'ont donc aucune question en commun ; les suivants complètent la base
jusqu'à ce que chaque question soit sortie au moins une fois (35 lots, limite fixée par le §7.8).

Dépendance : pip install openpyxl
Usage      : python3 outils/construire_lots.py   (depuis la racine du dépôt)
"""
import csv, glob, json, math, os, random, re
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
GRAINE = 2026

GRILLE = OrderedDict([
    ("1.1", 2), ("1.2.1", 2), ("1.2.2", 1), ("1.3", 3), ("1.5.1", 2), ("1.5.2", 2), ("1.8", 2),
    ("2.1", 2), ("2.2", 2), ("2.3", 2),
    ("3", 3), ("4", 2),
    ("5.1", 4), ("5.2", 2),
    ("6.1.1", 5), ("6.1.2", 2), ("6.2", 5), ("6.3", 3), ("6.4", 3), ("6.5", 2), ("6.6", 1),
    ("6.7", 1), ("6.8", 2), ("6.9", 1),
    ("7.1", 3), ("7.2", 3), ("7.3", 3), ("7.4", 2), ("7.5", 1), ("7.6", 2), ("7.7", 1),
    ("7.8", 2), ("7.9", 2), ("7.10", 2),
    ("8.1", 1), ("8.2.1", 1), ("8.2.2", 2), ("8.4", 2), ("8.5", 3), ("8.6", 1), ("8.7", 15),
    ("9.1", 2), ("9.2", 2), ("9.3", 1), ("9.4", 1), ("9.5", 1),
    ("10.1", 2), ("10.2", 1),
    ("11.1", 1), ("11.2", 1),
    ("12.1", 2), ("12.2", 1), ("12.3", 2), ("12.4", 1),
])
CAT_A = {"2.1", "2.2", "2.3", "3", "4", "5.1", "6.1.1", "6.2", "6.3", "6.4", "9.2"}
assert sum(GRILLE.values()) == 120 and sum(GRILLE[s] for s in CAT_A) == 33

# Réponses qui renvoient aux autres propositions : on les affiche pour que la réponse ait un sens
CONTEXTE = re.compile(
    r"(les deux (réponses|propositions|affirmations|notions|catégories)|dans les deux cas|^\s*les deux\b"
    r"|aucune des (réponses|propositions|affirmations|solutions)|ni l'une ni l'autre|ni l'un ni l'autre"
    r"|^\s*(a|b|c|1|2|3)\s*(,|et)\s*(a|b|c|1|2|3)\b|\bréponses? [abc]\b|ci-dessus|précédentes?\b)", re.I)
# Questions « négatives » : la réponse est l'élément qui ne convient PAS
NEGATIVE = re.compile(
    r"((lequel|laquelle|lesquels|lesquelles|quel|quelle|quels|quelles|parmi)\b[^?]{0,40}?\b"
    r"(n'est pas|ne sont pas|ne constitue pas|ne constituent pas|n'a pas|n'ont pas|ne fait pas partie"
    r"|ne font pas partie|n'entre pas|ne relève pas|ne figure pas|ne peut pas|ne peuvent pas|ne doit pas"
    r"|n'existe pas|ne s'applique pas|ne correspond pas|ne concerne pas|ne permet pas)|\b(est|sont) (fausses?|faux|inexactes?|incorrectes?|erronées?)\b"
    r"|(affirmations?|propositions?) (fausse|inexacte|incorrecte)|\bsauf\s*[:?]?\s*$|sauf une)", re.I)
FLAG = {"err": "Erreur de la base", "abs": "Réponse absente de la base", "obs": "Réponse dépassée",
        "ver": "À vérifier", "dis": "Formulation discutable", "pie": "Piège", "ctx": "Contexte"}

# --- Base et intitulés ---------------------------------------------------------
base = json.load(open("donnees/questions.json", encoding="utf-8"))
themes, labels = [], {"3": "Sécurité financière (blanchiment, corruption, embargos)", "4": "Abus de marché"}
for path in sorted(glob.glob("manuel/*.md")):
    txt = open(path, encoding="utf-8").read()
    n = int(os.path.basename(path)[:2])
    themes.append(dict(n=n, title=re.search(r"^# Thème \d+ — (.+)$", txt, re.M).group(1).strip()))
    for num, lab in re.findall(r"^## (\d+(?:\.\d+)+) (.+)$", txt, re.M):
        if num in GRILLE and num not in labels:
            labels[num] = re.sub(r"\s*\((catégorie A|\d+ questions)\)", "", lab).replace("*", "").strip()
assert set(labels) == set(GRILLE)


def num_key(uid):
    m = re.match(r"(\d+)([ab]?)", uid)
    return int(m.group(1)), m.group(2)


questions, by_sub = [], {}
for d in base:
    choix = [d["choix_a"], d["choix_b"], d["choix_c"]]
    a = "ABC".index(d["bonne_reponse"])
    q = OrderedDict(i=d["numero"], t=d["theme"], s=d["sous_theme"], k=d["categorie_examen"],
                    q=d["question"], r=choix[a])
    if CONTEXTE.search(choix[a]):
        q["c"] = [f"{'ABC'[j]}. {c}" for j, c in enumerate(choix)]
    if NEGATIVE.search(d["question"]):
        q["g"] = 1
    if d["alerte"]:
        q["f"], q["n"] = d["alerte"], d["note"]
    questions.append(q)
    by_sub.setdefault(q["s"], []).append(q["i"])
assert len(questions) == 2244 and set(by_sub) == set(GRILLE)

# --- Lots ------------------------------------------------------------------------
rng = random.Random(GRAINE)
NB_LOTS = max(math.ceil(len(by_sub[s]) / n) for s, n in GRILLE.items())
lots = [[] for _ in range(NB_LOTS)]
for s, n in GRILLE.items():
    ids = sorted(by_sub[s], key=num_key)
    rng.shuffle(ids)
    for k in range(NB_LOTS):
        lots[k] += [ids[(k * n + j) % len(ids)] for j in range(n)]

first = {}
for k, lot in enumerate(lots):
    assert len(lot) == len(set(lot)) == 120
    for u in lot:
        first.setdefault(u, k)
assert len(first) == 2244
fresh = [sum(1 for u in lot if first[u] == k) for k, lot in enumerate(lots)]
sans_rep = next(k for k, f in enumerate(fresh) if f < 120)

meta = dict(themes=themes, subs=[dict(s=s, t=int(s.split(".")[0]), label=labels[s], n=n,
                                      k="A" if s in CAT_A else "C", bank=len(by_sub[s])) for s, n in GRILLE.items()],
            seed=GRAINE, nb=NB_LOTS, sansRepetition=sans_rep)
os.makedirs("docs/data", exist_ok=True)
with open("docs/data/lots-data.js", "w", encoding="utf-8") as f:
    f.write("/* Lots de 120 questions : généré par outils/construire_lots.py, ne pas modifier à la main */\n")
    for name, obj in (("AMF_META", meta), ("AMF_Q", questions), ("AMF_LOTS", lots)):
        f.write(f"window.{name}=" + json.dumps(obj, ensure_ascii=False, separators=(",", ":")) + ";\n")

# --- Exports ---------------------------------------------------------------------
Q = {q["i"]: q for q in questions}
os.makedirs("donnees/lots", exist_ok=True)
head = ["Lot", "N° dans le lot", "Thème", "Sous-thème", "Catégorie", "N° question", "Question",
        "Bonne réponse", "Première apparition", "Alerte", "Note"]


def row(k, j, u):
    q = Q[u]
    rep = q["r"] + ("\n" + "\n".join(q["c"]) if "c" in q else "")
    return [k + 1, j + 1, q["t"], q["s"], q["k"], u, q["q"], rep,
            "nouvelle" if first[u] == k else f"lot {first[u] + 1}", FLAG.get(q.get("f"), ""), q.get("n", "")]


with open("donnees/lots/lots_120_questions.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(head)
    for k, lot in enumerate(lots):
        for j, u in enumerate(lot):
            w.writerow(row(k, j, u))

wb = Workbook()
ws = wb.active
ws.title = "Sommaire"
bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
fill = PatternFill("solid", fgColor="FFF2CC")
ws.append(["Lot", "Questions", "Catégorie A", "Catégorie C", "Nouvelles questions", "Déjà sorties dans un lot précédent"])
for k, lot in enumerate(lots):
    a = sum(1 for u in lot if Q[u]["k"] == "A")
    ws.append([k + 1, len(lot), a, len(lot) - a, fresh[k], len(lot) - fresh[k]])
ws.append([])
ws.append([f"Lots 1 à {sans_rep} : aucune question en commun. Lots suivants : complètent la base "
           f"jusqu'à ce que les 2 244 questions soient sorties au moins une fois."])
for c in ws[1]:
    c.font = bold
for col, wdt in zip("ABCDEF", (8, 11, 13, 13, 20, 32)):
    ws.column_dimensions[col].width = wdt
for k, lot in enumerate(lots):
    sh = wb.create_sheet(f"Lot {k + 1:02d}")
    sh.append(head[1:])
    for c in sh[1]:
        c.font = bold
    for j, u in enumerate(lot):
        sh.append(row(k, j, u)[1:])
        for c in sh[sh.max_row]:
            c.alignment = wrap
        sh.cell(sh.max_row, 7).fill = fill
    for col, wdt in zip("ABCDEFGHIJ", (8, 7, 10, 10, 11, 70, 50, 14, 20, 50)):
        sh.column_dimensions[col].width = wdt
    sh.freeze_panes = "A2"
wb.save("donnees/lots/Lots_120_questions.xlsx")

print(f"{NB_LOTS} lots de 120 questions ; lots 1 à {sans_rep} sans répétition")
print("nouvelles questions par lot :", fresh)
